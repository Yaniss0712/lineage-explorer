"""Construit l'architecture de test FICTIVE : fichiers Excel (1 par table) + JSON pour Lineage Explorer.

Architecture : Lac de données (10 tables brutes) → BV (8 tables de base) → Mart (5 tables métier).
Chaque table contient au plus 5 lignes ; les valeurs des BV et des marts sont réellement
calculées à partir des tables sources (jointures, renommages, agrégats), donc cohérentes.

Usage : python "architecture test/construire_architecture.py"   (nécessite openpyxl)

Produit, dans ce dossier :
- lac/, bv/, mart/            : un fichier Excel par table (onglets « donnees » + « dictionnaire ») ;
- 00_vue_ensemble.xlsx        : tables, flux (lineage), jointures, lineage colonne ;
- metadata_architecture.json  : JSON complet (structure + flux déclarés + relations) ;
- metadata_structure_seule.json : structure seule (+ clés) pour tester l'inférence ;
- verite_terrain.json         : les vrais flux, pour `python -m engine evaluate`.
"""

from __future__ import annotations

import json
import os
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
LOAD = "2026-10-01 08:00:00"
LAYER_DIR = {"LAC": "lac", "BV": "bv", "MART": "mart"}
TABLE_STYLE = {"LAC": "TableStyleMedium2", "BV": "TableStyleMedium7", "MART": "TableStyleMedium4"}

TABLES: dict[str, dict] = {}       # id → définition
FLOWS: list[dict] = []
RELATIONS: list[dict] = []


def table(tid, layer, ttype, description, columns, rows):
    """Déclare une table. columns = [(nom, type, est_cle, description)]."""
    database, name = tid.split(".")
    TABLES[tid] = {"id": tid, "database": database, "name": name, "layer": layer, "type": ttype,
                   "description": description, "columns": columns,
                   "rows": [dict(zip([c[0] for c in columns], r)) for r in rows]}
    assert len(rows) <= 5, tid
    return TABLES[tid]["rows"]


def flow(source, target, job, mapping):
    """Flux source → cible. mapping = [(col_source, col_cible, transformation | None)]."""
    FLOWS.append({"source": source, "target": target, "job": job,
                  "column_mapping": [{"source": s, "target": t, "transformation": tr} for s, t, tr in mapping]})


def relation(a, a_col, b, b_col, comment):
    RELATIONS.append({"a": a, "a_column": a_col, "b": b, "b_column": b_col, "comment": comment})


# ============================================================================ LAC
def mensualite(montant, taux, mois):
    """Mensualité d'un prêt amortissable (taux annuel, durée en mois)."""
    r = taux / 12
    return round(montant * r / (1 - (1 + r) ** -mois), 2)


S, I, D, DT, DEC = "string", "int", "date", "timestamp", "decimal(18,2)"

clients = table("lac_crm.clients_raw", "LAC", "TABLE", "Clients bruts issus du CRM", [
    ("code_tiers", S, True, "Identifiant tiers CRM"), ("nom_client", S, False, "Nom ou raison sociale"),
    ("segment", S, False, "Segment CRM (PART, PRO, ENT)"), ("code_agence", S, False, "Agence de rattachement"),
    ("date_creation", D, False, "Date d'entrée en relation"), ("date_chargement", DT, False, "Horodatage technique")], [
    (" t001", "Durand SARL", "PRO", "A01", "2019-03-12", LOAD),
    ("t002", "Martin Paul", "PART", "A02", "2020-07-01", LOAD),
    ("T003 ", "Bernard & Fils", "PRO", "A01", "2018-11-23", LOAD),
    ("t004", "Petit Claire", "PART", "A03", "2021-02-15", LOAD),
    ("T005", "Lefèvre Industrie", "ENT", "A02", "2017-06-30", LOAD)])

adresses = table("lac_crm.adresses_raw", "LAC", "TABLE", "Adresses postales des tiers", [
    ("id_adresse", S, True, "Identifiant adresse"), ("code_tiers", S, False, "Tiers propriétaire de l'adresse"),
    ("ville", S, False, "Ville"), ("code_postal", S, False, "Code postal"), ("pays_iso", S, False, "Pays (ISO 2)"),
    ("date_chargement", DT, False, "Horodatage technique")], [
    ("AD01", " t001", "Lyon", "69002", "FR", LOAD), ("AD02", "t002", "Nantes", "44000", "FR", LOAD),
    ("AD03", "T003 ", "Lyon", "69007", "FR", LOAD), ("AD04", "t004", "Lille", "59000", "FR", LOAD),
    ("AD05", "T005", "Toulouse", "31000", "FR", LOAD)])

agences = table("lac_ref.agences_raw", "LAC", "TABLE", "Référentiel des agences", [
    ("code_agence", S, True, "Code agence"), ("libelle_agence", S, False, "Nom de l'agence"),
    ("code_region", S, False, "Région de l'agence"), ("date_chargement", DT, False, "Horodatage technique")], [
    ("A01", "Lyon Centre", "R1", LOAD), ("A02", "Nantes Gare", "R2", LOAD), ("A03", "Lille Grand Place", "R3", LOAD)])

regions = table("lac_ref.regions_raw", "LAC", "TABLE", "Référentiel des régions", [
    ("code_region", S, True, "Code région"), ("nom_region", S, False, "Nom de la région"),
    ("date_chargement", DT, False, "Horodatage technique")], [
    ("R1", "Auvergne-Rhône-Alpes", LOAD), ("R2", "Pays de la Loire", LOAD), ("R3", "Hauts-de-France", LOAD)])

produits = table("lac_ref.produits_raw", "LAC", "TABLE", "Catalogue des produits de crédit", [
    ("code_produit", S, True, "Code produit"), ("libelle_produit", S, False, "Libellé produit"),
    ("famille_produit", S, False, "Famille (IMMO, CONSO, PRO)"), ("taux_base", "decimal(9,4)", False, "Taux annuel de base"),
    ("date_chargement", DT, False, "Horodatage technique")], [
    ("P10", "Prêt immobilier", "IMMO", 0.0325, LOAD), ("P20", "Prêt auto", "CONSO", 0.0480, LOAD),
    ("P30", "Prêt travaux", "CONSO", 0.0420, LOAD), ("P40", "Crédit professionnel", "PRO", 0.0390, LOAD)])

contrats = table("lac_pret.contrats_raw", "LAC", "TABLE", "Contrats de prêt bruts", [
    ("num_contrat", S, True, "Numéro de contrat"), ("code_tiers", S, False, "Emprunteur"),
    ("code_produit", S, False, "Produit souscrit"), ("code_agence", S, False, "Agence gestionnaire"),
    ("montant", DEC, False, "Montant emprunté"), ("date_signature", D, False, "Date de signature"),
    ("duree_mois", I, False, "Durée en mois"), ("date_chargement", DT, False, "Horodatage technique")], [
    ("C1001", " t001", "P40", "A01", 120000.00, "2023-04-10", 60, LOAD),
    ("C1002", "t002", "P10", "A02", 210000.00, "2022-09-01", 240, LOAD),
    ("C1003", "t002", "P20", "A02", 18000.00, "2024-01-15", 48, LOAD),
    ("C1004", "t004", "P30", "A03", 25000.00, "2024-06-20", 84, LOAD),
    ("C1005", "T005", "P40", "A02", 350000.00, "2021-12-05", 84, LOAD)])

echeances = table("lac_pret.echeances_raw", "LAC", "TABLE", "Échéances de remboursement", [
    ("id_echeance", S, True, "Identifiant échéance"), ("num_contrat", S, False, "Contrat concerné"),
    ("date_echeance", D, False, "Date d'échéance"), ("montant_echeance", DEC, False, "Montant dû"),
    ("statut_paiement", S, False, "PAYE / IMPAYE"), ("date_chargement", DT, False, "Horodatage technique")], [
    ("E01", "C1001", "2026-09-10", mensualite(120000, 0.039, 60), "PAYE", LOAD), ("E02", "C1002", "2026-09-01", mensualite(210000, 0.0325, 240), "PAYE", LOAD),
    ("E03", "C1003", "2026-09-15", mensualite(18000, 0.048, 48), "IMPAYE", LOAD), ("E04", "C1004", "2026-09-20", mensualite(25000, 0.042, 84), "IMPAYE", LOAD),
    ("E05", "C1005", "2026-09-05", mensualite(350000, 0.039, 84), "PAYE", LOAD)])

comptes = table("lac_cpt.comptes_raw", "LAC", "TABLE", "Comptes bancaires", [
    ("num_compte", S, True, "Numéro de compte"), ("code_tiers", S, False, "Titulaire"),
    ("code_agence", S, False, "Agence de tenue de compte"), ("type_compte", S, False, "Type de compte"),
    ("date_ouverture", D, False, "Date d'ouverture"), ("date_chargement", DT, False, "Horodatage technique")], [
    ("CPT01", " t001", "A01", "COURANT_PRO", "2019-03-12", LOAD), ("CPT02", "t002", "A02", "COURANT", "2020-07-01", LOAD),
    ("CPT03", "T003 ", "A01", "COURANT_PRO", "2018-11-23", LOAD), ("CPT04", "t004", "A03", "EPARGNE", "2021-02-15", LOAD),
    ("CPT05", "T005", "A02", "COURANT_PRO", "2017-06-30", LOAD)])

operations = table("lac_cpt.operations_raw", "LAC", "TABLE", "Opérations sur les comptes", [
    ("id_operation", S, True, "Identifiant opération"), ("num_compte", S, False, "Compte mouvementé"),
    ("date_operation", D, False, "Date de l'opération"), ("montant_op", DEC, False, "Montant (positif)"),
    ("sens", S, False, "C = crédit, D = débit"), ("code_canal", S, False, "Canal utilisé"),
    ("date_chargement", DT, False, "Horodatage technique")], [
    ("OP01", "CPT01", "2026-09-03", 5400.00, "C", "WEB", LOAD), ("OP02", "CPT01", "2026-09-08", 1200.00, "D", "GAB", LOAD),
    ("OP03", "CPT02", "2026-09-12", 2300.00, "C", "AGENCE", LOAD), ("OP04", "CPT03", "2026-09-18", 860.00, "D", "WEB", LOAD),
    ("OP05", "CPT05", "2026-09-25", 15000.00, "C", "WEB", LOAD)])

canaux = table("lac_ref.canaux_raw", "LAC", "TABLE", "Référentiel des canaux", [
    ("code_canal", S, True, "Code canal"), ("libelle_canal", S, False, "Libellé du canal"),
    ("date_chargement", DT, False, "Horodatage technique")], [
    ("WEB", "Banque en ligne", LOAD), ("GAB", "Distributeur", LOAD), ("AGENCE", "Guichet agence", LOAD)])


def by(rows, key):
    return {r[key]: r for r in rows}


def norm_tiers(code):
    return code.strip().upper()


# ============================================================================ BV
adr_by_tiers = by(adresses, "code_tiers")
tiers = table("bv_client.tiers", "BV", "VIEW", "Référentiel tiers nettoyé (identifiant normalisé, adresse)", [
    ("code_tiers_ref", S, True, "Identifiant tiers normalisé (trim + upper)"), ("nom_client", S, False, "Nom du client"),
    ("categorie_client", S, False, "Catégorie client (ex-segment CRM)"), ("code_agence", S, False, "Agence de rattachement"),
    ("ville", S, False, "Ville"), ("code_postal", S, False, "Code postal"), ("date_creation", D, False, "Entrée en relation")],
    [(norm_tiers(c["code_tiers"]), c["nom_client"], c["segment"], c["code_agence"],
      adr_by_tiers[c["code_tiers"]]["ville"], adr_by_tiers[c["code_tiers"]]["code_postal"], c["date_creation"])
     for c in clients])
flow("lac_crm.clients_raw", "bv_client.tiers", "job_bv_tiers", [
    ("code_tiers", "code_tiers_ref", "upper(trim(code_tiers))"), ("nom_client", "nom_client", None),
    ("segment", "categorie_client", "renommage"), ("code_agence", "code_agence", None),
    ("date_creation", "date_creation", None)])
flow("lac_crm.adresses_raw", "bv_client.tiers", "job_bv_tiers", [
    ("ville", "ville", None), ("code_postal", "code_postal", None)])

reg_by = by(regions, "code_region")
bv_agences = table("bv_ref.agences", "BV", "VIEW", "Agences avec leur région", [
    ("code_agence", S, True, "Code agence"), ("libelle_agence", S, False, "Nom de l'agence"),
    ("code_region", S, False, "Code région"), ("nom_region", S, False, "Nom de la région")],
    [(a["code_agence"], a["libelle_agence"], a["code_region"], reg_by[a["code_region"]]["nom_region"]) for a in agences])
flow("lac_ref.agences_raw", "bv_ref.agences", "job_bv_ref", [
    ("code_agence", "code_agence", None), ("libelle_agence", "libelle_agence", None), ("code_region", "code_region", None)])
flow("lac_ref.regions_raw", "bv_ref.agences", "job_bv_ref", [("nom_region", "nom_region", None)])

bv_produits = table("bv_ref.produits", "BV", "VIEW", "Catalogue produits", [
    ("code_produit", S, True, "Code produit"), ("libelle_produit", S, False, "Libellé produit"),
    ("famille_produit", S, False, "Famille de produit"), ("taux_base", "decimal(9,4)", False, "Taux annuel de base")],
    [(p["code_produit"], p["libelle_produit"], p["famille_produit"], p["taux_base"]) for p in produits])
flow("lac_ref.produits_raw", "bv_ref.produits", "job_bv_ref", [
    (c, c, None) for c in ("code_produit", "libelle_produit", "famille_produit", "taux_base")])

bv_contrats = table("bv_pret.contrats", "BV", "VIEW", "Contrats de prêt nettoyés", [
    ("num_contrat", S, True, "Numéro de contrat"), ("code_tiers_ref", S, False, "Emprunteur (identifiant normalisé)"),
    ("code_produit", S, False, "Produit"), ("code_agence", S, False, "Agence gestionnaire"),
    ("montant_pret", DEC, False, "Montant emprunté"), ("date_signature", D, False, "Date de signature"),
    ("duree_mois", I, False, "Durée en mois")],
    [(c["num_contrat"], norm_tiers(c["code_tiers"]), c["code_produit"], c["code_agence"], c["montant"],
      c["date_signature"], c["duree_mois"]) for c in contrats])
flow("lac_pret.contrats_raw", "bv_pret.contrats", "job_bv_pret", [
    ("num_contrat", "num_contrat", None), ("code_tiers", "code_tiers_ref", "upper(trim(code_tiers))"),
    ("code_produit", "code_produit", None), ("code_agence", "code_agence", None), ("montant", "montant_pret", "renommage"),
    ("date_signature", "date_signature", None), ("duree_mois", "duree_mois", None)])

bv_echeances = table("bv_pret.echeances", "BV", "VIEW", "Échéances de remboursement", [
    ("id_echeance", S, True, "Identifiant échéance"), ("num_contrat", S, False, "Contrat"),
    ("date_echeance", D, False, "Date d'échéance"), ("montant_echeance", DEC, False, "Montant dû"),
    ("statut_paiement", S, False, "PAYE / IMPAYE")],
    [(e["id_echeance"], e["num_contrat"], e["date_echeance"], e["montant_echeance"], e["statut_paiement"]) for e in echeances])
flow("lac_pret.echeances_raw", "bv_pret.echeances", "job_bv_pret", [
    (c, c, None) for c in ("id_echeance", "num_contrat", "date_echeance", "montant_echeance", "statut_paiement")])

bv_comptes = table("bv_cpt.comptes", "BV", "VIEW", "Comptes bancaires nettoyés", [
    ("num_compte", S, True, "Numéro de compte"), ("code_tiers_ref", S, False, "Titulaire (identifiant normalisé)"),
    ("code_agence", S, False, "Agence"), ("type_compte", S, False, "Type de compte"), ("date_ouverture", D, False, "Ouverture")],
    [(c["num_compte"], norm_tiers(c["code_tiers"]), c["code_agence"], c["type_compte"], c["date_ouverture"]) for c in comptes])
flow("lac_cpt.comptes_raw", "bv_cpt.comptes", "job_bv_cpt", [
    ("num_compte", "num_compte", None), ("code_tiers", "code_tiers_ref", "upper(trim(code_tiers))"),
    ("code_agence", "code_agence", None), ("type_compte", "type_compte", None), ("date_ouverture", "date_ouverture", None)])

canal_by = by(canaux, "code_canal")
bv_operations = table("bv_cpt.operations", "BV", "VIEW", "Opérations avec libellé de canal", [
    ("id_operation", S, True, "Identifiant opération"), ("num_compte", S, False, "Compte"),
    ("date_operation", D, False, "Date"), ("montant_operation", DEC, False, "Montant (positif)"),
    ("sens", S, False, "C / D"), ("code_canal", S, False, "Canal"), ("libelle_canal", S, False, "Libellé du canal")],
    [(o["id_operation"], o["num_compte"], o["date_operation"], o["montant_op"], o["sens"], o["code_canal"],
      canal_by[o["code_canal"]]["libelle_canal"]) for o in operations])
flow("lac_cpt.operations_raw", "bv_cpt.operations", "job_bv_cpt", [
    ("id_operation", "id_operation", None), ("num_compte", "num_compte", None), ("date_operation", "date_operation", None),
    ("montant_op", "montant_operation", "renommage"), ("sens", "sens", None), ("code_canal", "code_canal", None)])
flow("lac_ref.canaux_raw", "bv_cpt.operations", "job_bv_cpt", [("libelle_canal", "libelle_canal", None)])

prod_by = by(bv_produits, "code_produit")


contrats_enrichis = table("bv_pret.contrats_enrichis", "BV", "VIEW", "Contrats enrichis du produit (BV construite sur des BV)", [
    ("num_contrat", S, True, "Numéro de contrat"), ("code_tiers_ref", S, False, "Emprunteur"),
    ("code_agence", S, False, "Agence"), ("code_produit", S, False, "Produit"), ("libelle_produit", S, False, "Libellé produit"),
    ("famille_produit", S, False, "Famille"), ("taux_base", "decimal(9,4)", False, "Taux annuel"),
    ("montant_pret", DEC, False, "Montant emprunté"), ("duree_mois", I, False, "Durée"),
    ("date_signature", D, False, "Signature"), ("mensualite_theorique", DEC, False, "Mensualité calculée")],
    [(c["num_contrat"], c["code_tiers_ref"], c["code_agence"], c["code_produit"],
      prod_by[c["code_produit"]]["libelle_produit"], prod_by[c["code_produit"]]["famille_produit"],
      prod_by[c["code_produit"]]["taux_base"], c["montant_pret"], c["duree_mois"], c["date_signature"],
      mensualite(c["montant_pret"], prod_by[c["code_produit"]]["taux_base"], c["duree_mois"])) for c in bv_contrats])
flow("bv_pret.contrats", "bv_pret.contrats_enrichis", "job_bv_contrats_enrichis", [
    (c, c, None) for c in ("num_contrat", "code_tiers_ref", "code_agence", "code_produit", "montant_pret",
                           "duree_mois", "date_signature")]
    + [("montant_pret", "mensualite_theorique", "montant_pret * (taux_base/12) / (1 - (1 + taux_base/12)^-duree_mois)")])
flow("bv_ref.produits", "bv_pret.contrats_enrichis", "job_bv_contrats_enrichis", [
    ("libelle_produit", "libelle_produit", None), ("famille_produit", "famille_produit", None),
    ("taux_base", "taux_base", None)])

# ============================================================================ MART
tiers_by = by(tiers, "code_tiers_ref")
ag_by = by(bv_agences, "code_agence")
ctr_by_tiers = defaultdict(list)
for c in bv_contrats:
    ctr_by_tiers[c["code_tiers_ref"]].append(c)
cpt_by_tiers = Counter(c["code_tiers_ref"] for c in bv_comptes)

table("dm_client.vue_360", "MART", "TABLE", "Vue 360° du client", [
    ("code_tiers_ref", S, True, "Client"), ("nom_client", S, False, "Nom"), ("categorie_client", S, False, "Catégorie"),
    ("ville", S, False, "Ville"), ("libelle_agence", S, False, "Agence"), ("nb_comptes", I, False, "Nombre de comptes"),
    ("nb_contrats", I, False, "Nombre de prêts"), ("montant_total_pret", DEC, False, "Total emprunté")],
    [(t["code_tiers_ref"], t["nom_client"], t["categorie_client"], t["ville"], ag_by[t["code_agence"]]["libelle_agence"],
      cpt_by_tiers[t["code_tiers_ref"]], len(ctr_by_tiers[t["code_tiers_ref"]]),
      sum(c["montant_pret"] for c in ctr_by_tiers[t["code_tiers_ref"]])) for t in tiers])
flow("bv_client.tiers", "dm_client.vue_360", "job_dm_vue_360", [
    (c, c, None) for c in ("code_tiers_ref", "nom_client", "categorie_client", "ville")])
flow("bv_ref.agences", "dm_client.vue_360", "job_dm_vue_360", [("libelle_agence", "libelle_agence", "jointure sur code_agence")])
flow("bv_cpt.comptes", "dm_client.vue_360", "job_dm_vue_360", [("num_compte", "nb_comptes", "count(num_compte)")])
flow("bv_pret.contrats", "dm_client.vue_360", "job_dm_vue_360", [
    ("num_contrat", "nb_contrats", "count(num_contrat)"), ("montant_pret", "montant_total_pret", "sum(montant_pret)")])

ech_by_ctr = defaultdict(list)
for e in bv_echeances:
    ech_by_ctr[e["num_contrat"]].append(e)


def impaye(contracts):
    return round(sum(e["montant_echeance"] for c in contracts for e in ech_by_ctr[c["num_contrat"]]
                     if e["statut_paiement"] == "IMPAYE"), 2)


table("dm_pret.encours_client", "MART", "TABLE", "Encours de prêt par client", [
    ("code_tiers_ref", S, True, "Client"), ("nom_client", S, False, "Nom"), ("categorie_client", S, False, "Catégorie"),
    ("nb_contrats", I, False, "Nombre de prêts"), ("montant_total_pret", DEC, False, "Total emprunté"),
    ("montant_impaye", DEC, False, "Total des échéances impayées")],
    [(t, tiers_by[t]["nom_client"], tiers_by[t]["categorie_client"], len(cs), sum(c["montant_pret"] for c in cs), impaye(cs))
     for t, cs in sorted(ctr_by_tiers.items())])
flow("bv_client.tiers", "dm_pret.encours_client", "job_dm_encours", [
    (c, c, None) for c in ("code_tiers_ref", "nom_client", "categorie_client")])
flow("bv_pret.contrats", "dm_pret.encours_client", "job_dm_encours", [
    ("num_contrat", "nb_contrats", "count(num_contrat)"), ("montant_pret", "montant_total_pret", "sum(montant_pret)")])
flow("bv_pret.echeances", "dm_pret.encours_client", "job_dm_encours", [
    ("montant_echeance", "montant_impaye", "sum(montant_echeance) where statut_paiement = 'IMPAYE'")])

enr_by = by(contrats_enrichis, "num_contrat")
table("dm_pret.impayes", "MART", "TABLE", "Échéances impayées détaillées", [
    ("num_contrat", S, True, "Contrat"), ("code_tiers_ref", S, False, "Client"), ("libelle_produit", S, False, "Produit"),
    ("famille_produit", S, False, "Famille"), ("date_echeance", D, False, "Échéance"), ("montant_impaye", DEC, False, "Montant impayé")],
    [(e["num_contrat"], enr_by[e["num_contrat"]]["code_tiers_ref"], enr_by[e["num_contrat"]]["libelle_produit"],
      enr_by[e["num_contrat"]]["famille_produit"], e["date_echeance"], e["montant_echeance"])
     for e in bv_echeances if e["statut_paiement"] == "IMPAYE"])
flow("bv_pret.echeances", "dm_pret.impayes", "job_dm_impayes", [
    ("num_contrat", "num_contrat", None), ("date_echeance", "date_echeance", None),
    ("montant_echeance", "montant_impaye", "filtre statut_paiement = 'IMPAYE'")])
flow("bv_pret.contrats_enrichis", "dm_pret.impayes", "job_dm_impayes", [
    (c, c, None) for c in ("code_tiers_ref", "libelle_produit", "famille_produit")])

table("dm_reseau.activite_agence", "MART", "TABLE", "Activité commerciale par agence", [
    ("code_agence", S, True, "Agence"), ("libelle_agence", S, False, "Nom"), ("nom_region", S, False, "Région"),
    ("nb_clients", I, False, "Clients rattachés"), ("nb_contrats", I, False, "Prêts gérés"),
    ("nb_comptes", I, False, "Comptes tenus"), ("montant_total_pret", DEC, False, "Total emprunté")],
    [(a["code_agence"], a["libelle_agence"], a["nom_region"],
      sum(1 for t in tiers if t["code_agence"] == a["code_agence"]),
      sum(1 for c in bv_contrats if c["code_agence"] == a["code_agence"]),
      sum(1 for c in bv_comptes if c["code_agence"] == a["code_agence"]),
      sum(c["montant_pret"] for c in bv_contrats if c["code_agence"] == a["code_agence"])) for a in bv_agences])
flow("bv_ref.agences", "dm_reseau.activite_agence", "job_dm_agence", [
    (c, c, None) for c in ("code_agence", "libelle_agence", "nom_region")])
flow("bv_client.tiers", "dm_reseau.activite_agence", "job_dm_agence", [("code_tiers_ref", "nb_clients", "count(code_tiers_ref)")])
flow("bv_pret.contrats", "dm_reseau.activite_agence", "job_dm_agence", [
    ("num_contrat", "nb_contrats", "count(num_contrat)"), ("montant_pret", "montant_total_pret", "sum(montant_pret)")])
flow("bv_cpt.comptes", "dm_reseau.activite_agence", "job_dm_agence", [("num_compte", "nb_comptes", "count(num_compte)")])

cpt_by = by(bv_comptes, "num_compte")
ops_by_cpt = defaultdict(list)
for o in bv_operations:
    ops_by_cpt[o["num_compte"]].append(o)
table("dm_cpt.flux_mensuels", "MART", "TABLE", "Flux mensuels par compte (septembre 2026)", [
    ("num_compte", S, True, "Compte"), ("code_tiers_ref", S, False, "Titulaire"), ("categorie_client", S, False, "Catégorie"),
    ("mois", S, False, "Mois (AAAA-MM)"), ("total_credits", DEC, False, "Somme des crédits"),
    ("total_debits", DEC, False, "Somme des débits"), ("nb_operations", I, False, "Nombre d'opérations"),
    ("canal_principal", S, False, "Canal le plus utilisé")],
    [(n, cpt_by[n]["code_tiers_ref"], tiers_by[cpt_by[n]["code_tiers_ref"]]["categorie_client"], "2026-09",
      sum(o["montant_operation"] for o in ops if o["sens"] == "C"),
      sum(o["montant_operation"] for o in ops if o["sens"] == "D"), len(ops),
      Counter(o["libelle_canal"] for o in ops).most_common(1)[0][0]) for n, ops in sorted(ops_by_cpt.items())])
flow("bv_cpt.operations", "dm_cpt.flux_mensuels", "job_dm_flux", [
    ("num_compte", "num_compte", None), ("date_operation", "mois", "substr(date_operation, 1, 7)"),
    ("montant_operation", "total_credits", "sum(montant_operation) where sens = 'C'"),
    ("montant_operation", "total_debits", "sum(montant_operation) where sens = 'D'"),
    ("id_operation", "nb_operations", "count(id_operation)"), ("libelle_canal", "canal_principal", "mode(libelle_canal)")])
flow("bv_cpt.comptes", "dm_cpt.flux_mensuels", "job_dm_flux", [("code_tiers_ref", "code_tiers_ref", None)])
flow("bv_client.tiers", "dm_cpt.flux_mensuels", "job_dm_flux", [("categorie_client", "categorie_client", "jointure sur code_tiers_ref")])

# ============================================================================ JOINTURES
JOIN_KEYS = [  # (colonne de la table de référence, [(table, colonne) qui s'y joignent], commentaire)
    (("lac_crm.clients_raw", "code_tiers"), [("lac_crm.adresses_raw", "code_tiers"), ("lac_pret.contrats_raw", "code_tiers"),
                                             ("lac_cpt.comptes_raw", "code_tiers")], "Tiers (identifiant brut)"),
    (("bv_client.tiers", "code_tiers_ref"), [("bv_pret.contrats", "code_tiers_ref"), ("bv_cpt.comptes", "code_tiers_ref"),
                                             ("bv_pret.contrats_enrichis", "code_tiers_ref"), ("dm_client.vue_360", "code_tiers_ref"),
                                             ("dm_pret.encours_client", "code_tiers_ref"), ("dm_cpt.flux_mensuels", "code_tiers_ref")],
     "Tiers (identifiant normalisé)"),
    (("lac_ref.agences_raw", "code_agence"), [("lac_crm.clients_raw", "code_agence"), ("lac_pret.contrats_raw", "code_agence"),
                                              ("lac_cpt.comptes_raw", "code_agence")], "Agence"),
    (("bv_ref.agences", "code_agence"), [("bv_client.tiers", "code_agence"), ("bv_pret.contrats", "code_agence"),
                                         ("bv_cpt.comptes", "code_agence"), ("dm_reseau.activite_agence", "code_agence")], "Agence"),
    (("lac_ref.regions_raw", "code_region"), [("lac_ref.agences_raw", "code_region")], "Région"),
    (("lac_ref.produits_raw", "code_produit"), [("lac_pret.contrats_raw", "code_produit")], "Produit"),
    (("bv_ref.produits", "code_produit"), [("bv_pret.contrats", "code_produit"), ("bv_pret.contrats_enrichis", "code_produit")], "Produit"),
    (("lac_pret.contrats_raw", "num_contrat"), [("lac_pret.echeances_raw", "num_contrat")], "Contrat"),
    (("bv_pret.contrats", "num_contrat"), [("bv_pret.echeances", "num_contrat")], "Contrat"),
    (("bv_pret.contrats_enrichis", "num_contrat"), [("bv_pret.echeances", "num_contrat"), ("dm_pret.impayes", "num_contrat")], "Contrat"),
    (("lac_cpt.comptes_raw", "num_compte"), [("lac_cpt.operations_raw", "num_compte")], "Compte"),
    (("bv_cpt.comptes", "num_compte"), [("bv_cpt.operations", "num_compte"), ("dm_cpt.flux_mensuels", "num_compte")], "Compte"),
    (("lac_ref.canaux_raw", "code_canal"), [("lac_cpt.operations_raw", "code_canal")], "Canal"),
]
flow_pairs = {tuple(sorted((f["source"], f["target"]))) for f in FLOWS}
for (ref, ref_col), others, comment in JOIN_KEYS:
    for other, other_col in others:
        if tuple(sorted((ref, other))) not in flow_pairs:   # un flux explique déjà le lien
            relation(ref, ref_col, other, other_col, f"Jointure {comment} : {ref}.{ref_col} = {other}.{other_col}")


# ============================================================================ ÉCRITURE EXCEL
ARIAL = Font(name="Arial", size=10)
ARIAL_B = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=13, bold=True)
NUM_FORMATS = {DEC: "#,##0.00", "decimal(9,4)": "0.00%", I: "0"}


def autosize(ws):
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(60, max(10, width + 3))


def add_table(ws, name, headers, rows, style, start_row=1, formats=None):
    for j, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=j, value=h).font = ARIAL_B
    for i, row in enumerate(rows, start_row + 1):
        for j, value in enumerate(row, 1):
            cell = ws.cell(row=i, column=j, value=value)
            cell.font = ARIAL
            if formats and formats[j - 1]:
                cell.number_format = formats[j - 1]
    ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + max(1, len(rows))}"
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def sources_of(tid):
    """Colonne cible → liste « source.colonne (transformation) »."""
    result = defaultdict(list)
    for f in FLOWS:
        if f["target"] == tid:
            for m in f["column_mapping"]:
                result[m["target"]].append(f"{f['source']}.{m['source']}" + (f" ({m['transformation']})" if m["transformation"] else ""))
    return result


def write_table_file(t):
    wb = Workbook()
    ws = wb.active
    ws.title = "donnees"
    cols = t["columns"]
    add_table(ws, "T_" + t["id"].replace(".", "_"), [c[0] for c in cols],
              [[r[c[0]] for c in cols] for r in t["rows"]], TABLE_STYLE[t["layer"]],
              formats=[NUM_FORMATS.get(c[1]) for c in cols])
    autosize(ws)
    wd = wb.create_sheet("dictionnaire")
    wd["A1"] = f"{t['id']} — couche {t['layer']} ({t['type']})"
    wd["A1"].font = TITLE
    wd["A2"] = t["description"] + ". Données FICTIVES ; valeurs des BV / marts calculées à partir des sources ci-dessous."
    wd["A2"].font = ARIAL
    srcs = sources_of(t["id"])
    add_table(wd, "D_" + t["id"].replace(".", "_"), ["colonne", "type", "cle", "description", "alimentee_par"],
              [[c[0], c[1], "oui" if c[2] else "", c[3], " ; ".join(srcs.get(c[0], [])) or ("—" if t["layer"] == "LAC" else "")]
               for c in cols], "TableStyleLight9", start_row=4)
    autosize(wd)
    wd.column_dimensions["A"].width = 24
    folder = os.path.join(HERE, LAYER_DIR[t["layer"]])
    os.makedirs(folder, exist_ok=True)
    wb.save(os.path.join(folder, f"{t['id']}.xlsx"))


def write_overview():
    wb = Workbook()
    ws = wb.active
    ws.title = "lisez_moi"
    lines = ["Architecture de test FICTIVE — Lineage Explorer", "",
             f"{sum(1 for t in TABLES.values() if t['layer'] == 'LAC')} tables du lac → "
             f"{sum(1 for t in TABLES.values() if t['layer'] == 'BV')} tables BV → "
             f"{sum(1 for t in TABLES.values() if t['layer'] == 'MART')} marts ; {len(FLOWS)} flux ; {len(RELATIONS)} jointures.",
             "Onglet tables : liste des tables · flux : qui alimente qui · jointures : clés de jointure · "
             "lineage_colonnes : origine de chaque colonne.",
             "Les données (≤ 5 lignes par table) sont dans les dossiers lac/, bv/, mart/ (un fichier par table).",
             "JSON pour l'outil : metadata_architecture.json (complet) ou metadata_structure_seule.json (test de l'inférence)."]
    for i, line in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=line).font = TITLE if i == 1 else ARIAL
    ws.column_dimensions["A"].width = 120

    wt = wb.create_sheet("tables")
    add_table(wt, "Tables", ["couche", "table", "type", "nb_colonnes", "nb_lignes", "description", "fichier"],
              [[t["layer"], t["id"], t["type"], len(t["columns"]), len(t["rows"]), t["description"],
                f"{LAYER_DIR[t['layer']]}/{t['id']}.xlsx"] for t in TABLES.values()], "TableStyleMedium2")
    autosize(wt)

    wf = wb.create_sheet("flux")
    add_table(wf, "Flux", ["source", "couche_source", "cible", "couche_cible", "job", "nb_colonnes"],
              [[f["source"], TABLES[f["source"]]["layer"], f["target"], TABLES[f["target"]]["layer"], f["job"],
                len(f["column_mapping"])] for f in FLOWS], "TableStyleMedium7")
    autosize(wf)

    wr = wb.create_sheet("jointures")
    add_table(wr, "Jointures", ["table_a", "colonne_a", "table_b", "colonne_b", "commentaire"],
              [[r["a"], r["a_column"], r["b"], r["b_column"], r["comment"]] for r in RELATIONS], "TableStyleMedium4")
    autosize(wr)

    wl = wb.create_sheet("lineage_colonnes")
    add_table(wl, "LineageColonnes", ["source", "colonne_source", "cible", "colonne_cible", "transformation"],
              [[f["source"], m["source"], f["target"], m["target"], m["transformation"] or "copie"]
               for f in FLOWS for m in f["column_mapping"]], "TableStyleLight9")
    autosize(wl)
    wb.save(os.path.join(HERE, "00_vue_ensemble.xlsx"))


# ============================================================================ ÉCRITURE JSON
def json_objects(with_details: bool) -> list[dict]:
    objects = []
    for t in TABLES.values():
        obj = {"id": t["id"], "database": t["database"], "name": t["name"], "type": t["type"],
               "description": t["description"],
               "columns": [{"name": c[0], "type": c[1], "is_key": c[2], "description": c[3], "position": i}
                           for i, c in enumerate(t["columns"], 1)]}
        if with_details:
            obj["layer"] = t["layer"]
            obj["owner"] = "equipe_data_fictive"
        objects.append(obj)
    return objects


def write_json():
    meta = {"format_version": "1.0", "extracted_at": "2026-10-09T10:00:00+02:00",
            "source": "architecture-test-fictive", "extractor": "construire_architecture.py"}
    full = {"meta": meta, "objects": json_objects(True), "flows": FLOWS, "relations": RELATIONS}
    structure = {"meta": {**meta, "source": "architecture-test-fictive (structure seule)"}, "objects": json_objects(False)}
    truth = {"comment": "Vrais flux de l'architecture de test (fictive)",
             "flows": [{"source": f["source"], "target": f["target"]} for f in FLOWS]}
    for name, data in (("metadata_architecture.json", full), ("metadata_structure_seule.json", structure),
                       ("verite_terrain.json", truth)):
        with open(os.path.join(HERE, name), "w", encoding="utf-8") as handle:
            json.dump(data, handle, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    for t in TABLES.values():
        write_table_file(t)
    write_overview()
    write_json()
    layers = Counter(t["layer"] for t in TABLES.values())
    print(f"{len(TABLES)} tables ({dict(layers)}), {len(FLOWS)} flux, {len(RELATIONS)} jointures")
